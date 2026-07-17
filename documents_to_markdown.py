#!/usr/bin/env python3
"""
Document to Markdown converter.

Extracts readable text from `.pdf`, `.epub`, `.docx`, `.doc`, and `.pptx` files and
builds structured Markdown files that are easier for AI tools to ingest.
Supports single-file and whole-directory processing with both CLI and tkinter
GUI entry points.
"""

from __future__ import annotations

import argparse
import re
import shutil
import statistics
import subprocess
import sys
import tempfile
import threading
import zipfile
from dataclasses import dataclass
from html import unescape
from html.parser import HTMLParser
from pathlib import Path
from queue import Empty, Queue
from xml.etree import ElementTree as ET

import pdfplumber
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {".pdf", ".epub", ".docx", ".doc", ".pptx"}
LIST_RE = re.compile(r"^(?P<marker>(?:[-*\u2022o])|(?:\d+[.)]))\s+(?P<text>.+)$")
ROMAN_RE = re.compile(r"^(?=[ivxlcdmIVXLCDM]+$)[IVXLCDMivxlcdm]{1,8}$")
WORD_NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
}
EPUB_NS = {
    "container": "urn:oasis:names:tc:opendocument:xmlns:container",
    "opf": "http://www.idpf.org/2007/opf",
}


@dataclass
class PdfLine:
    text: str
    size: float
    x0: float
    top: float
    bottom: float
    page_number: int
    bold: bool
    monospace: bool


@dataclass
class MarkdownBlock:
    kind: str
    text: str = ""
    level: int = 0
    ordered: bool = False


@dataclass
class ConversionIssue:
    source: Path
    error: str


@dataclass
class SlideBlock:
    kind: str
    text: str = ""
    level: int = 0
    rows: list[list[str]] | None = None


@dataclass
class SlideContent:
    title: str
    blocks: list[SlideBlock]


class HtmlToMarkdownParser(HTMLParser):
    HEADING_TAGS = {"h1", "h2", "h3", "h4", "h5", "h6"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.blocks: list[MarkdownBlock] = []
        self.stack: list[str] = []
        self.current_tag: str | None = None
        self.current_text: list[str] = []
        self.list_depth = 0

    def handle_starttag(self, tag, attrs):
        tag = tag.lower()
        self.stack.append(tag)
        if tag in {"p", "blockquote"} and "li" in self.stack and self.current_tag == "li":
            return
        if tag in self.HEADING_TAGS or tag in {"p", "li", "pre", "blockquote"}:
            self._flush_current()
            self.current_tag = tag
            self.current_text = []
        elif tag in {"ul", "ol"}:
            self._flush_current()
            self.list_depth += 1
        elif tag == "br":
            self.current_text.append("\n")

    def handle_endtag(self, tag):
        tag = tag.lower()
        if tag == self.current_tag:
            self._flush_current()
        elif tag in {"ul", "ol"}:
            self._flush_current()
            self.list_depth = max(0, self.list_depth - 1)
        for index in range(len(self.stack) - 1, -1, -1):
            if self.stack[index] == tag:
                del self.stack[index]
                break

    def handle_data(self, data):
        if self.current_tag:
            self.current_text.append(data)

    def close(self):
        super().close()
        self._flush_current()

    def _flush_current(self):
        if not self.current_tag:
            return

        text = clean_text(" ".join(part.strip() for part in self.current_text if part.strip()))
        if text:
            if self.current_tag in self.HEADING_TAGS:
                level = int(self.current_tag[1])
                self.blocks.append(MarkdownBlock(kind="heading", text=text, level=max(2, min(level + 1, 4))))
            elif self.current_tag == "li":
                self.blocks.append(MarkdownBlock(kind="list", text=text, level=max(0, self.list_depth - 1)))
            elif self.current_tag == "pre":
                self.blocks.append(MarkdownBlock(kind="code", text=text))
            else:
                self.blocks.append(MarkdownBlock(kind="paragraph", text=text))

        self.current_tag = None
        self.current_text = []


def clean_text(value: str) -> str:
    return " ".join((value or "").replace("\u00a0", " ").split())


def escape_markdown_cell(value: str) -> str:
    return clean_text(value).replace("|", "\\|")


def iter_powerpoint_text_shapes(shapes):
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_powerpoint_text_shapes(shape.shapes)
            continue
        if getattr(shape, "has_table", False) or getattr(shape, "has_text_frame", False):
            yield shape


def powerpoint_placeholder_type(shape):
    if not getattr(shape, "is_placeholder", False):
        return None
    try:
        return shape.placeholder_format.type
    except (AttributeError, ValueError):
        return None


def is_powerpoint_title(shape) -> bool:
    return powerpoint_placeholder_type(shape) in {
        PP_PLACEHOLDER.TITLE,
        PP_PLACEHOLDER.CENTER_TITLE,
    }


def is_powerpoint_subtitle(shape) -> bool:
    return powerpoint_placeholder_type(shape) == PP_PLACEHOLDER.SUBTITLE


def powerpoint_paragraph_text(paragraph) -> str:
    runs = "".join(run.text for run in paragraph.runs) if paragraph.runs else paragraph.text
    return clean_text(runs)


def powerpoint_font_sizes(shape) -> list[float]:
    sizes: list[float] = []
    if not getattr(shape, "has_text_frame", False):
        return sizes
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            size = getattr(getattr(run, "font", None), "size", None)
            if size is not None:
                sizes.append(float(size.pt))
    return sizes


def powerpoint_paragraph_is_bold(paragraph) -> bool:
    bold_flags = [
        bool(run.font.bold)
        for run in paragraph.runs
        if getattr(getattr(run, "font", None), "bold", None) is not None
    ]
    return bool(bold_flags) and all(bold_flags)


def powerpoint_paragraph_is_heading(
    text: str,
    paragraph,
    base_font_size: float | None,
    shape,
) -> bool:
    if not text or len(text) > 80 or text.endswith((".", "!", "?", ";")):
        return False
    if getattr(paragraph, "level", 0):
        return False
    if is_powerpoint_subtitle(shape):
        return True

    sizes = [
        float(run.font.size.pt)
        for run in paragraph.runs
        if getattr(getattr(run, "font", None), "size", None) is not None
    ]
    paragraph_size = max(sizes) if sizes else None
    words = [word for word in text.split() if any(character.isalpha() for character in word)]
    title_like = bool(words) and sum(word[:1].isupper() for word in words) >= max(
        1, len(words) // 2
    )
    return bool(
        (powerpoint_paragraph_is_bold(paragraph) and title_like)
        or (
            paragraph_size
            and base_font_size
            and paragraph_size >= base_font_size * 1.2
            and title_like
        )
    )


def extract_powerpoint_shape_blocks(shape, base_font_size: float | None) -> list[SlideBlock]:
    blocks: list[SlideBlock] = []
    if getattr(shape, "has_table", False):
        rows = [
            [escape_markdown_cell(cell.text) for cell in row.cells]
            for row in shape.table.rows
        ]
        rows = [row for row in rows if any(row)]
        return [SlideBlock(kind="table", rows=rows)] if rows else []

    if not getattr(shape, "has_text_frame", False):
        return blocks
    for paragraph in shape.text_frame.paragraphs:
        text = powerpoint_paragraph_text(paragraph)
        if not text:
            continue
        level = max(0, int(getattr(paragraph, "level", 0) or 0))
        if powerpoint_paragraph_is_heading(text, paragraph, base_font_size, shape):
            blocks.append(SlideBlock(kind="heading", text=text, level=3))
        elif level > 0:
            blocks.append(SlideBlock(kind="list", text=text, level=level))
        else:
            blocks.append(SlideBlock(kind="paragraph", text=text))
    return blocks


def extract_slide_content(slide, slide_number: int) -> SlideContent:
    title = ""
    blocks: list[SlideBlock] = []
    shapes = list(iter_powerpoint_text_shapes(slide.shapes))
    font_sizes = [size for shape in shapes for size in powerpoint_font_sizes(shape)]
    base_font_size = statistics.median(font_sizes) if font_sizes else None

    for shape in shapes:
        shape_blocks = extract_powerpoint_shape_blocks(shape, base_font_size)
        if not shape_blocks:
            continue
        if is_powerpoint_title(shape) and not title:
            title = shape_blocks[0].text
            shape_blocks = shape_blocks[1:]
        blocks.extend(shape_blocks)

    return SlideContent(title=title or f"Slide {slide_number}", blocks=blocks)


def powerpoint_title(presentation: Presentation, source_path: Path) -> str:
    core_title = clean_text(getattr(presentation.core_properties, "title", "") or "")
    if core_title:
        return core_title
    for index, slide in enumerate(presentation.slides, start=1):
        extracted = extract_slide_content(slide, index)
        if extracted.title and not extracted.title.startswith("Slide "):
            return extracted.title
    return source_path.stem.replace("_", " ").strip() or source_path.stem


def render_powerpoint_table(rows: list[list[str]]) -> list[str]:
    width = max(len(row) for row in rows)
    padded_rows = [row + [""] * (width - len(row)) for row in rows]
    lines = [
        "| " + " | ".join(padded_rows[0]) + " |",
        "| " + " | ".join(["---"] * width) + " |",
    ]
    lines.extend("| " + " | ".join(row) + " |" for row in padded_rows[1:])
    return lines


def render_slide_markdown(index: int, slide: SlideContent) -> list[str]:
    lines = [f"## {index:02d}. {slide.title}", ""]
    if not slide.blocks:
        return [*lines, "_No extractable text on this slide._", ""]

    paragraph_buffer: list[str] = []

    def flush_paragraph() -> None:
        if paragraph_buffer:
            lines.extend([" ".join(paragraph_buffer), ""])
            paragraph_buffer.clear()

    for block in slide.blocks:
        if block.kind == "paragraph":
            paragraph_buffer.append(block.text)
            continue
        flush_paragraph()
        if block.kind == "heading":
            lines.extend([f"### {block.text}", ""])
        elif block.kind == "list":
            lines.append(f"{'  ' * max(0, block.level - 1)}- {block.text}")
        elif block.kind == "table" and block.rows:
            lines.extend([*render_powerpoint_table(block.rows), ""])
    flush_paragraph()
    if lines[-1] != "":
        lines.append("")
    return lines


def metadata_title(reader: PdfReader) -> str:
    try:
        raw_title = getattr(reader.metadata, "title", "") or ""
    except Exception:
        raw_title = ""
    return clean_text(raw_title)


def looks_like_heading(line: PdfLine, body_size: float) -> bool:
    text = line.text
    if not text or len(text) > 100:
        return False
    if text.endswith((".", "!", "?", ";")):
        return False

    words = [word for word in text.split() if any(ch.isalpha() for ch in word)]
    if not words:
        return False

    titled = sum(1 for word in words if word[:1].isupper())
    title_like = titled >= max(1, len(words) // 2) or text.isupper()
    return title_like and (line.size >= body_size * 1.18 or line.bold)


def join_text(parts: list[str], next_text: str) -> None:
    if not parts:
        parts.append(next_text)
        return
    previous = parts[-1]
    if previous.endswith("-") and next_text[:1].islower():
        parts[-1] = previous[:-1] + next_text
    else:
        parts[-1] = previous + " " + next_text


def group_words_into_lines(page, page_number: int) -> list[PdfLine]:
    words = page.extract_words(
        x_tolerance=2,
        y_tolerance=3,
        use_text_flow=True,
        extra_attrs=["size", "fontname"],
    )
    if not words:
        return []

    words = sorted(words, key=lambda word: (round(word["top"], 1), word["x0"]))
    grouped: list[list[dict]] = []

    for word in words:
        if not grouped:
            grouped.append([word])
            continue

        current = grouped[-1]
        current_top = statistics.mean(item["top"] for item in current)
        tolerance = max(2.5, float(word.get("size", 10)) * 0.35)
        if abs(word["top"] - current_top) <= tolerance:
            current.append(word)
        else:
            grouped.append([word])

    lines: list[PdfLine] = []
    for group in grouped:
        text = clean_text(" ".join(item["text"] for item in sorted(group, key=lambda value: value["x0"])))
        if not text:
            continue

        sizes = [float(item.get("size", 10)) for item in group]
        fonts = [str(item.get("fontname", "")) for item in group]
        line = PdfLine(
            text=text,
            size=max(sizes) if sizes else 10.0,
            x0=min(item["x0"] for item in group),
            top=min(item["top"] for item in group),
            bottom=max(item["bottom"] for item in group),
            page_number=page_number,
            bold=any("bold" in font.lower() for font in fonts),
            monospace=all(
                any(token in font.lower() for token in ("courier", "mono", "consolas", "menlo"))
                for font in fonts
            ),
        )
        lines.append(line)

    return lines


def extract_lines(pdf_path: Path) -> list[PdfLine]:
    lines: list[PdfLine] = []
    with pdfplumber.open(str(pdf_path)) as pdf:
        for page_number, page in enumerate(pdf.pages, start=1):
            lines.extend(group_words_into_lines(page, page_number))
    return lines


def line_is_noise(line: PdfLine) -> bool:
    text = line.text.strip()
    if not text:
        return True
    if re.fullmatch(r"\d{1,4}", text):
        return True
    if ROMAN_RE.fullmatch(text):
        return True
    return False


def heading_levels(lines: list[PdfLine], body_size: float) -> dict[float, int]:
    sizes = sorted(
        {
            round(line.size, 1)
            for line in lines
            if not line_is_noise(line) and looks_like_heading(line, body_size)
        },
        reverse=True,
    )
    levels: dict[float, int] = {}
    for index, size in enumerate(sizes[:3], start=2):
        levels[size] = index
    return levels


def infer_title(lines: list[PdfLine], fallback: str) -> str:
    page_one = [line for line in lines if line.page_number == 1 and not line_is_noise(line)]
    if not page_one:
        return fallback

    largest = max(page_one, key=lambda line: (line.size, -line.top))
    if largest.size >= statistics.median(line.size for line in page_one) * 1.2 and len(largest.text) <= 120:
        return largest.text
    return fallback


def lines_to_blocks(lines: list[PdfLine]) -> tuple[str, list[MarkdownBlock]]:
    meaningful = [line for line in lines if not line_is_noise(line)]
    if not meaningful:
        return "Untitled Document", []

    body_size = statistics.median(line.size for line in meaningful)
    title = infer_title(meaningful, "Untitled Document")
    level_map = heading_levels(meaningful, body_size)
    left_margin = statistics.median(line.x0 for line in meaningful)

    blocks: list[MarkdownBlock] = []
    paragraph_parts: list[str] = []
    code_lines: list[str] = []
    previous_line: PdfLine | None = None

    def flush_paragraph():
        if paragraph_parts:
            blocks.append(MarkdownBlock(kind="paragraph", text=paragraph_parts[0]))
            paragraph_parts.clear()

    def flush_code():
        if code_lines:
            blocks.append(MarkdownBlock(kind="code", text="\n".join(code_lines)))
            code_lines.clear()

    for line in meaningful:
        if line.text == title and line.page_number == 1:
            previous_line = line
            continue

        list_match = LIST_RE.match(line.text)
        if list_match:
            flush_paragraph()
            flush_code()
            indent = max(0, round((line.x0 - left_margin) / 18))
            blocks.append(
                MarkdownBlock(
                    kind="list",
                    text=list_match.group("text"),
                    level=indent,
                    ordered=list_match.group("marker")[0].isdigit(),
                )
            )
            previous_line = line
            continue

        if line.monospace and (line.x0 - left_margin) > 8:
            flush_paragraph()
            code_lines.append(line.text)
            previous_line = line
            continue

        flush_code()

        rounded_size = round(line.size, 1)
        if rounded_size in level_map and looks_like_heading(line, body_size):
            flush_paragraph()
            blocks.append(MarkdownBlock(kind="heading", text=line.text, level=level_map[rounded_size]))
            previous_line = line
            continue

        new_paragraph = False
        if previous_line is None:
            new_paragraph = True
        elif line.page_number != previous_line.page_number:
            new_paragraph = True
        elif (line.top - previous_line.bottom) > max(4, line.size * 0.9):
            new_paragraph = True
        elif abs(line.x0 - previous_line.x0) > 20:
            new_paragraph = True

        if new_paragraph:
            flush_paragraph()
            paragraph_parts.append(line.text)
        else:
            join_text(paragraph_parts, line.text)

        previous_line = line

    flush_paragraph()
    flush_code()
    return title, blocks


def render_markdown(title: str, blocks: list[MarkdownBlock]) -> str:
    lines = [f"# {title}", ""]

    for block in blocks:
        if block.kind == "heading":
            lines.extend([f"{'#' * min(max(block.level, 2), 4)} {block.text}", ""])
        elif block.kind == "paragraph":
            lines.extend([block.text, ""])
        elif block.kind == "code":
            lines.extend(["```", block.text, "```", ""])
        elif block.kind == "list":
            indent = "  " * block.level
            marker = "1." if block.ordered else "-"
            lines.append(f"{indent}{marker} {block.text}")

    if lines and lines[-1] != "":
        lines.append("")
    return "\n".join(lines).rstrip() + "\n"


class DocumentToMarkdownConverter:
    supported_extensions = SUPPORTED_EXTENSIONS

    def extract_blocks(self, pdf_path: Path) -> tuple[str, list[MarkdownBlock]]:
        reader = PdfReader(str(pdf_path))
        fallback_title = metadata_title(reader) or pdf_path.stem.replace("_", " ").strip() or pdf_path.stem
        lines = extract_lines(pdf_path)
        title, blocks = lines_to_blocks(lines)
        if title == "Untitled Document":
            title = fallback_title
        return title, blocks

    def extract_powerpoint_slides(self, pptx_path: Path) -> tuple[str, list[SlideContent]]:
        presentation = Presentation(str(pptx_path))
        title = powerpoint_title(presentation, pptx_path)
        slides = [
            extract_slide_content(slide, index)
            for index, slide in enumerate(presentation.slides, start=1)
        ]
        return title, slides

    def extract_docx_blocks(self, docx_path: Path) -> tuple[str, list[MarkdownBlock]]:
        paragraphs: list[MarkdownBlock] = []
        title = docx_path.stem.replace("_", " ").strip() or docx_path.stem

        with zipfile.ZipFile(docx_path) as archive:
            document_xml = archive.read("word/document.xml")

        root = ET.fromstring(document_xml)
        body = root.find("w:body", WORD_NS)
        if body is None:
            return title, []

        for child in body:
            if child.tag == f"{{{WORD_NS['w']}}}p":
                block = self._docx_paragraph_to_block(child)
                if not block:
                    continue
                if title == docx_path.stem.replace("_", " ").strip() and block.kind == "heading" and block.level == 2:
                    title = block.text
                    continue
                paragraphs.append(block)
            elif child.tag == f"{{{WORD_NS['w']}}}tbl":
                rows = self._docx_table_rows(child)
                if rows:
                    paragraphs.append(MarkdownBlock(kind="paragraph", text="\n".join(rows)))

        return title, paragraphs

    def _docx_paragraph_to_block(self, paragraph) -> MarkdownBlock | None:
        text = clean_text("".join(node.text or "" for node in paragraph.findall(".//w:t", WORD_NS)))
        if not text:
            return None

        style = ""
        p_style = paragraph.find("w:pPr/w:pStyle", WORD_NS)
        if p_style is not None:
            style = p_style.attrib.get(f"{{{WORD_NS['w']}}}val", "")

        num_pr = paragraph.find("w:pPr/w:numPr", WORD_NS)
        if num_pr is not None or style.lower().startswith("list"):
            ilvl = paragraph.find("w:pPr/w:numPr/w:ilvl", WORD_NS)
            level = 0
            if ilvl is not None:
                try:
                    level = int(ilvl.attrib.get(f"{{{WORD_NS['w']}}}val", "0"))
                except ValueError:
                    level = 0
            return MarkdownBlock(kind="list", text=text, level=level)

        heading_match = re.match(r"heading([1-6])", style.replace(" ", "").lower())
        if heading_match:
            return MarkdownBlock(kind="heading", text=text, level=max(2, min(int(heading_match.group(1)) + 1, 4)))

        return MarkdownBlock(kind="paragraph", text=text)

    def _docx_table_rows(self, table) -> list[str]:
        rows: list[str] = []
        for row in table.findall("w:tr", WORD_NS):
            cells = []
            for cell in row.findall("w:tc", WORD_NS):
                text = clean_text(" ".join(node.text or "" for node in cell.findall(".//w:t", WORD_NS)))
                cells.append(text)
            if cells:
                rows.append("| " + " | ".join(cells) + " |")
        return rows

    def extract_epub_blocks(self, epub_path: Path) -> tuple[str, list[MarkdownBlock]]:
        with zipfile.ZipFile(epub_path) as archive:
            opf_path = self._epub_opf_path(archive)
            opf_root = ET.fromstring(archive.read(opf_path))
            base = Path(opf_path).parent
            manifest = {
                item.attrib["id"]: item.attrib
                for item in opf_root.findall(".//opf:manifest/opf:item", EPUB_NS)
                if "id" in item.attrib and "href" in item.attrib
            }
            title_node = opf_root.find(".//{http://purl.org/dc/elements/1.1/}title")
            title = clean_text(title_node.text if title_node is not None else "") or epub_path.stem.replace("_", " ")
            blocks: list[MarkdownBlock] = []

            for itemref in opf_root.findall(".//opf:spine/opf:itemref", EPUB_NS):
                item = manifest.get(itemref.attrib.get("idref", ""))
                if not item:
                    continue
                media_type = item.get("media-type", "")
                if media_type not in {"application/xhtml+xml", "text/html"}:
                    continue
                content_path = (base / item["href"]).as_posix()
                parser = HtmlToMarkdownParser()
                parser.feed(unescape(archive.read(content_path).decode("utf-8", errors="replace")))
                parser.close()
                blocks.extend(parser.blocks)

        return title, blocks

    def _epub_opf_path(self, archive: zipfile.ZipFile) -> str:
        container = ET.fromstring(archive.read("META-INF/container.xml"))
        rootfile = container.find(".//container:rootfile", EPUB_NS)
        if rootfile is None:
            raise ValueError("EPUB container does not point to an OPF package.")
        return rootfile.attrib["full-path"]

    def extract_doc_blocks(self, doc_path: Path) -> tuple[str, list[MarkdownBlock]]:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            converted = self._convert_legacy_doc_to_docx(doc_path, temp_path)
            return self.extract_docx_blocks(converted)

    def _convert_legacy_doc_to_docx(self, doc_path: Path, output_dir: Path) -> Path:
        soffice = shutil.which("soffice") or shutil.which("libreoffice")
        if soffice:
            subprocess.run(
                [
                    soffice,
                    "--headless",
                    "--convert-to",
                    "docx",
                    "--outdir",
                    str(output_dir),
                    str(doc_path),
                ],
                check=True,
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
                text=True,
            )
            converted = output_dir / f"{doc_path.stem}.docx"
            if converted.exists():
                return converted

        try:
            import win32com.client  # type: ignore
        except ImportError as exc:
            raise RuntimeError(
                "Legacy .doc conversion requires LibreOffice on PATH or Microsoft Word with pywin32 available."
            ) from exc

        converted = output_dir / f"{doc_path.stem}.docx"
        word = win32com.client.Dispatch("Word.Application")
        word.Visible = False
        document = None
        try:
            document = word.Documents.Open(str(doc_path.resolve()))
            document.SaveAs(str(converted.resolve()), FileFormat=16)
        finally:
            if document is not None:
                document.Close(False)
            word.Quit()

        if not converted.exists():
            raise RuntimeError("Legacy .doc conversion did not produce a .docx file.")
        return converted

    def extract_document_blocks(self, source_path: Path) -> tuple[str, list[MarkdownBlock]]:
        suffix = source_path.suffix.lower()
        if suffix == ".pdf":
            return self.extract_blocks(source_path)
        if suffix == ".docx":
            return self.extract_docx_blocks(source_path)
        if suffix == ".doc":
            return self.extract_doc_blocks(source_path)
        if suffix == ".epub":
            return self.extract_epub_blocks(source_path)
        raise ValueError(f"Unsupported file type: {source_path.suffix}")

    def build_markdown(self, source_path: Path, output_path: Path) -> Path:
        if source_path.suffix.lower() == ".pptx":
            title, slides = self.extract_powerpoint_slides(source_path)
            lines = [f"# {title}", ""]
            for index, slide in enumerate(slides, start=1):
                lines.extend(render_slide_markdown(index, slide))
            markdown = "\n".join(lines).rstrip() + "\n"
        else:
            title, blocks = self.extract_document_blocks(source_path)
            markdown = render_markdown(title, blocks)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_text(markdown, encoding="utf-8")
        return output_path

    def collect_inputs(self, input_path: Path) -> list[Path]:
        if input_path.is_file():
            if input_path.suffix.lower() not in self.supported_extensions:
                supported = ", ".join(sorted(self.supported_extensions))
                raise ValueError(f"Only {supported} files are supported.")
            return [input_path]

        if input_path.is_dir():
            files = sorted(
                (path for path in input_path.rglob("*") if path.is_file() and path.suffix.lower() in self.supported_extensions),
                key=lambda path: (path.stat().st_size, str(path).lower()),
            )
            if not files:
                supported = ", ".join(sorted(self.supported_extensions))
                raise ValueError(f"No supported files ({supported}) found in {input_path}")
            return files

        raise ValueError(f"Input path not found: {input_path}")

    def output_path_for(self, source_path: Path, input_path: Path, output_dir: Path, used_paths: set[Path]) -> Path:
        base_dir = input_path if input_path.is_dir() else input_path.parent
        if input_path.is_dir():
            relative_path = source_path.relative_to(base_dir).with_suffix(".md")
            output_path = output_dir / relative_path
        else:
            output_path = output_dir / f"{source_path.stem}.md"

        normalized = output_path.resolve()
        if normalized not in used_paths:
            used_paths.add(normalized)
            return output_path

        output_path = output_path.with_name(f"{source_path.stem}_{source_path.suffix.lower().lstrip('.')}.md")
        normalized = output_path.resolve()
        counter = 2
        while normalized in used_paths:
            output_path = output_path.with_name(f"{source_path.stem}_{source_path.suffix.lower().lstrip('.')}_{counter}.md")
            normalized = output_path.resolve()
            counter += 1
        used_paths.add(normalized)
        return output_path

    def convert(self, input_path: Path, output_dir: Path, progress_callback=None) -> list[Path]:
        files = self.collect_inputs(input_path)
        results: list[Path] = []
        failures: list[ConversionIssue] = []
        used_paths: set[Path] = set()

        if progress_callback:
            progress_callback(
                0,
                len(files),
                None,
                None,
                f"Found {len(files)} supported document file(s). Processing smaller files first.",
            )

        for index, source_path in enumerate(files, start=1):
            output_path = self.output_path_for(source_path, input_path, output_dir, used_paths)

            if progress_callback:
                size_mb = source_path.stat().st_size / (1024 * 1024)
                progress_callback(
                    index - 1,
                    len(files),
                    source_path,
                    output_path,
                    f"Processing {index}/{len(files)}: {source_path.name} ({size_mb:.1f} MB)",
                )

            try:
                output_path.parent.mkdir(parents=True, exist_ok=True)
                output_path.touch(exist_ok=True)
                result = self.build_markdown(source_path, output_path)
                results.append(result)
            except Exception as exc:
                failures.append(ConversionIssue(source=source_path, error=str(exc)))
                if progress_callback:
                    progress_callback(
                        index,
                        len(files),
                        source_path,
                        output_path,
                        f"Skipped {source_path.name}: {exc}",
                    )
                continue

            if progress_callback:
                progress_callback(index, len(files), source_path, result, None)

        if failures and progress_callback:
            progress_callback(
                len(results),
                len(files),
                None,
                None,
                f"Completed with {len(failures)} skipped file(s). Check the log for details.",
            )

        return results


PdfToMarkdownConverter = DocumentToMarkdownConverter


def launch_gui():
    import tkinter as tk
    from tkinter import filedialog, messagebox, scrolledtext, ttk

    from ui_icons import IconStore
    from ui_theme import (
        MONO_FONT,
        apply_ttk_theme,
        load_theme,
        save_theme,
        style_text_widget,
    )
    from window_geometry import centered_geometry

    class App:
        def __init__(self, root):
            self.root = root
            self.icons = IconStore()
            self._icon_targets = []
            self._status_kind = "info"
            self.icons.apply_window_icon(root, "app-documents-to-markdown")
            self.root.title("Documents to Markdown · Chwezi Document Suite")
            self.root.minsize(820, 650)
            self.theme = load_theme()
            self.palette = apply_ttk_theme(self.root, self.theme)

            self.mode = tk.StringVar(value="file")
            self.input_path = tk.StringVar()
            self.output_dir = tk.StringVar()
            self.status = tk.StringVar(value="Ready")
            self.progress = tk.DoubleVar(value=0.0)
            self.queue = Queue()
            self.worker = None

            self._build()
            self._apply_theme()
            self.root.after(100, self._drain_queue)

        def _build(self):
            main = ttk.Frame(self.root, padding=24, style="App.TFrame")
            main.grid(row=0, column=0, sticky="nsew")
            self.root.columnconfigure(0, weight=1)
            self.root.rowconfigure(0, weight=1)
            main.columnconfigure(0, weight=1)
            main.rowconfigure(6, weight=1)

            header = ttk.Frame(main, style="App.TFrame")
            header.grid(row=0, column=0, sticky="ew", pady=(0, 18))
            header.columnconfigure(0, weight=1)
            title_label = ttk.Label(
                header,
                text="Documents to Markdown",
                style="Title.TLabel",
            )
            title_label.grid(row=0, column=0, sticky="w")
            self._bind_icon(title_label, "apps", "app-documents-to-markdown", 30)
            ttk.Label(
                header,
                text="Extract useful structure from PDF, EPUB, Word, and PowerPoint files.",
                style="Subtitle.TLabel",
            ).grid(row=1, column=0, sticky="w", pady=(4, 0))
            self.theme_btn = ttk.Button(header, command=self._toggle_theme)
            self.theme_btn.grid(row=0, column=1, rowspan=2, sticky="e")

            mode_frame = ttk.LabelFrame(
                main,
                text=" 01 · Choose the workload ",
                style="Card.TLabelframe",
            )
            mode_frame.grid(row=1, column=0, sticky="ew", pady=(0, 12))
            mode_frame.columnconfigure((0, 1), weight=1)
            single_mode = ttk.Radiobutton(
                mode_frame,
                text="Single document",
                variable=self.mode,
                value="file",
                command=self._sync_defaults,
                style="Card.TRadiobutton",
            )
            single_mode.grid(row=0, column=0, sticky="w", padx=(0, 18), pady=2)
            self._bind_icon(single_mode, "actions", "file", 16)
            folder_mode = ttk.Radiobutton(
                mode_frame,
                text="A whole folder",
                variable=self.mode,
                value="directory",
                command=self._sync_defaults,
                style="Card.TRadiobutton",
            )
            folder_mode.grid(row=0, column=1, sticky="w", pady=2)
            self._bind_icon(folder_mode, "navigation", "folder-tree", 16)

            paths = ttk.LabelFrame(
                main,
                text=" 02 · Select source and destination ",
                style="Card.TLabelframe",
            )
            paths.grid(row=2, column=0, sticky="ew", pady=(0, 12))
            paths.columnconfigure(1, weight=1)

            self._file_row(
                paths,
                0,
                "Source",
                self.input_path,
                self._browse_input,
                ("navigation", "folder-open"),
            )
            self._file_row(
                paths,
                1,
                "Save Markdown in",
                self.output_dir,
                self._browse_output,
                ("actions", "folder-output"),
            )

            actions = ttk.Frame(main, style="App.TFrame")
            actions.grid(row=3, column=0, sticky="ew", pady=(2, 12))
            self.convert_btn = ttk.Button(
                actions,
                text="Convert to Markdown",
                command=self._start,
                style="Primary.TButton",
            )
            self.convert_btn.pack(side="left")
            self._bind_icon(self.convert_btn, "file-types", "file-type-markdown", 17)
            clear_button = ttk.Button(
                actions, text="Clear activity", command=self._clear_log
            )
            clear_button.pack(
                side="left", padx=(10, 0)
            )
            self._bind_icon(clear_button, "actions", "eraser", 16)

            self.status_label = ttk.Label(
                main, textvariable=self.status, style="Status.TLabel"
            )
            self.status_label.grid(
                row=4, column=0, sticky="ew", pady=(0, 8)
            )
            ttk.Progressbar(main, maximum=100, variable=self.progress).grid(
                row=5, column=0, sticky="ew", pady=(0, 12)
            )

            log_frame = ttk.LabelFrame(
                main,
                text=" Activity ",
                style="Card.TLabelframe",
            )
            log_frame.grid(row=6, column=0, sticky="nsew")
            log_frame.columnconfigure(0, weight=1)
            log_frame.rowconfigure(0, weight=1)
            self.log = scrolledtext.ScrolledText(
                log_frame,
                height=14,
                font=(MONO_FONT, 9),
                padx=12,
                pady=10,
                wrap="word",
            )
            self.log.grid(row=0, column=0, sticky="nsew")

        def _file_row(self, parent, row, label, variable, command, icon):
            ttk.Label(parent, text=label, style="Card.TLabel").grid(
                row=row, column=0, sticky="w", padx=(0, 12), pady=5
            )
            ttk.Entry(parent, textvariable=variable).grid(
                row=row, column=1, sticky="ew", padx=(0, 8), pady=5
            )
            button = ttk.Button(parent, text="Browse…", command=command)
            button.grid(row=row, column=2, pady=5)
            self._bind_icon(button, icon[0], icon[1], 15)

        def _bind_icon(self, widget, group, name, size, compound="left"):
            self._icon_targets.append((widget, group, name, size, compound))

        def _apply_icons(self):
            for widget, group, name, size, compound in self._icon_targets:
                image = self.icons.tk(group, name, size, self.theme)
                if image is not None:
                    widget.configure(image=image, compound=compound)
            theme_name = "sun" if self.theme == "dark" else "moon"
            theme_icon = self.icons.tk("navigation", theme_name, 15, self.theme)
            if theme_icon is not None:
                self.theme_btn.configure(image=theme_icon, compound="left")
            status_icons = {
                "info": ("status", "circle-info"),
                "loading": ("status", "loader-circle"),
                "success": ("status", "circle-check"),
                "error": ("status", "triangle-alert"),
            }
            group, name = status_icons[self._status_kind]
            status_icon = self.icons.tk(group, name, 16, self.theme)
            if status_icon is not None:
                self.status_label.configure(image=status_icon, compound="left")

        def _set_status(self, message, kind="info"):
            self.status.set(message)
            self._status_kind = kind
            self._apply_icons()

        def _apply_theme(self):
            self.palette = apply_ttk_theme(self.root, self.theme)
            style_text_widget(self.log, self.palette)
            self.theme_btn.configure(
                text="Light mode" if self.theme == "dark" else "Dark mode"
            )
            self._apply_icons()

        def _toggle_theme(self):
            self.theme = "light" if self.theme == "dark" else "dark"
            save_theme(self.theme)
            self._apply_theme()

        def _browse_input(self):
            if self.mode.get() == "directory":
                selected = filedialog.askdirectory(title="Select Document Directory")
            else:
                selected = filedialog.askopenfilename(
                    title="Select Document File",
                    filetypes=[
                        ("Supported documents", "*.pdf *.epub *.docx *.doc *.pptx"),
                        ("PDF files", "*.pdf"),
                        ("EPUB files", "*.epub"),
                        ("Word files", "*.docx *.doc"),
                        ("PowerPoint files", "*.pptx"),
                        ("All files", "*.*"),
                    ],
                )
            if selected:
                self.input_path.set(selected)
                self._sync_defaults()

        def _browse_output(self):
            selected = filedialog.askdirectory(title="Select Output Directory")
            if selected:
                self.output_dir.set(selected)

        def _sync_defaults(self):
            input_path = self.input_path.get().strip()
            if not input_path:
                return

            source = Path(input_path)
            if self.mode.get() == "directory":
                self.output_dir.set(str(source.parent / f"{source.name}_markdown"))
            else:
                self.output_dir.set(str(source.parent / "markdown"))

        def _append_log(self, message):
            self.log.insert("end", message + "\n")
            self.log.see("end")

        def _clear_log(self):
            self.log.delete("1.0", "end")

        def _start(self):
            input_value = self.input_path.get().strip()
            output_value = self.output_dir.get().strip()

            if not input_value:
                messagebox.showerror("Missing", "Please select a document file or directory.")
                return
            if not output_value:
                messagebox.showerror("Missing", "Please select an output directory.")
                return
            if self.worker and self.worker.is_alive():
                return

            input_path = Path(input_value)
            output_dir = Path(output_value)

            self._clear_log()
            self.progress.set(0)
            self._set_status("Starting conversion...", "loading")
            self._append_log(f"Input: {input_path}")
            self._append_log(f"Output: {output_dir}")
            self.convert_btn.state(["disabled"])

            self.worker = threading.Thread(
                target=self._run_worker,
                args=(input_path, output_dir),
                daemon=True,
            )
            self.worker.start()

        def _run_worker(self, input_path: Path, output_dir: Path):
            converter = DocumentToMarkdownConverter()

            def progress_callback(current, total, source_path, output_path, info_message):
                if info_message:
                    self.queue.put(("info", info_message))
                    return
                self.queue.put(
                    (
                        "progress",
                        (
                            current,
                            total,
                            f"Converted {source_path.name} -> {output_path.name}",
                        ),
                    )
                )

            try:
                results = converter.convert(
                    input_path,
                    output_dir,
                    progress_callback=progress_callback,
                )
                self.queue.put(("done", results))
            except Exception as exc:
                self.queue.put(("error", str(exc)))

        def _drain_queue(self):
            try:
                while True:
                    kind, payload = self.queue.get_nowait()
                    if kind == "info":
                        self._append_log(payload)
                        self._set_status(payload, "info")
                    elif kind == "progress":
                        current, total, message = payload
                        percent = 0 if total <= 0 else (current / total) * 100
                        self.progress.set(percent)
                        self._set_status(message, "loading")
                        self._append_log(message)
                    elif kind == "done":
                        self.progress.set(100)
                        message = f"Finished. Created {len(payload)} Markdown file(s)."
                        self._set_status(message, "success")
                        self._append_log(message)
                        self.convert_btn.state(["!disabled"])
                        messagebox.showinfo("Conversion Complete", message)
                    elif kind == "error":
                        self._set_status("Conversion failed.", "error")
                        self._append_log(f"ERROR: {payload}")
                        self.convert_btn.state(["!disabled"])
                        messagebox.showerror("Conversion Failed", payload)
            except Empty:
                pass
            finally:
                self.root.after(100, self._drain_queue)

    root = tk.Tk()
    App(root)
    root.update_idletasks()
    width = max(root.winfo_reqwidth(), 900)
    height = max(root.winfo_reqheight(), 690)
    root.geometry(
        centered_geometry(
            width,
            height,
            root.winfo_screenwidth(),
            root.winfo_screenheight(),
        )
    )
    root.mainloop()


def build_arg_parser():
    parser = argparse.ArgumentParser(
        description="Convert PDF, EPUB, Word, and PowerPoint files into structured Markdown files."
    )
    parser.add_argument("--gui", action="store_true", help="Launch the converter GUI.")
    parser.add_argument("--input", help="Path to a supported document file or a directory containing supported files.")
    parser.add_argument("--output-dir", help="Directory where Markdown files will be written.")
    return parser


def run_cli(args) -> int:
    if not args.input:
        raise SystemExit("--input is required in CLI mode")
    if not args.output_dir:
        raise SystemExit("--output-dir is required in CLI mode")

    converter = DocumentToMarkdownConverter()
    input_path = Path(args.input).expanduser()
    output_dir = Path(args.output_dir).expanduser()

    def progress_callback(current, total, source_path, output_path, info_message):
        if info_message:
            print(info_message)
            return
        print(f"[{current}/{total}] {source_path} -> {output_path}")

    results = converter.convert(input_path, output_dir, progress_callback=progress_callback)
    print(f"Created {len(results)} Markdown file(s) in {output_dir}")
    return 0


def main(argv=None) -> int:
    parser = build_arg_parser()
    argv = sys.argv[1:] if argv is None else argv

    if not argv or "--gui" in argv:
        launch_gui()
        return 0

    args = parser.parse_args(argv)
    return run_cli(args)


if __name__ == "__main__":
    raise SystemExit(main())
