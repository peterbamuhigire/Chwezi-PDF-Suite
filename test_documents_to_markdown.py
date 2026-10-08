"""
Smoke tests for the unified document-to-Markdown converter.

Run with: python test_documents_to_markdown.py
"""

import sys
import tempfile
import zipfile
from pathlib import Path

from pptx import Presentation
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

from documents_to_markdown import DocumentToMarkdownConverter


def create_sample_pdf(target: Path) -> Path:
    pdf = canvas.Canvas(str(target), pagesize=letter)
    pdf.setTitle("Sample PDF Book")
    pdf.setFont("Helvetica-Bold", 18)
    pdf.drawString(72, 720, "Sample PDF Book")
    pdf.setFont("Helvetica-Bold", 14)
    pdf.drawString(72, 690, "Getting Started")
    pdf.setFont("Helvetica", 12)
    pdf.drawString(72, 666, "This is the first paragraph on page one.")
    pdf.drawString(72, 650, "It should appear in the Markdown output.")
    pdf.drawString(72, 626, "1. First numbered item")
    pdf.drawString(72, 610, "2. Second numbered item")
    pdf.save()
    return target


def create_sample_docx(target: Path) -> Path:
    content_types = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml" ContentType="application/xml"/>
  <Override PartName="/word/document.xml"
    ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
</Types>"""
    rels = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1"
    Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument"
    Target="word/document.xml"/>
</Relationships>"""
    document = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:body>
    <w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>Sample DOCX Book</w:t></w:r></w:p>
    <w:p><w:r><w:t>This paragraph came from a Word document.</w:t></w:r></w:p>
    <w:p><w:pPr><w:pStyle w:val="ListParagraph"/></w:pPr><w:r><w:t>DOCX list item</w:t></w:r></w:p>
  </w:body>
</w:document>"""
    with zipfile.ZipFile(target, "w") as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr("_rels/.rels", rels)
        archive.writestr("word/document.xml", document)
    return target


def create_sample_epub(target: Path) -> Path:
    container = """<?xml version="1.0" encoding="UTF-8"?>
<container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">
  <rootfiles>
    <rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/>
  </rootfiles>
</container>"""
    opf = """<?xml version="1.0" encoding="UTF-8"?>
<package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="bookid">
  <metadata xmlns:dc="http://purl.org/dc/elements/1.1/">
    <dc:title>Sample EPUB Book</dc:title>
  </metadata>
  <manifest>
    <item id="chapter1" href="chapter1.xhtml" media-type="application/xhtml+xml"/>
  </manifest>
  <spine><itemref idref="chapter1"/></spine>
</package>"""
    chapter = """<?xml version="1.0" encoding="UTF-8"?>
<html xmlns="http://www.w3.org/1999/xhtml"><body>
  <h1>Opening Chapter</h1>
  <p>This paragraph came from an EPUB file.</p>
  <ul><li>EPUB list item</li></ul>
</body></html>"""
    with zipfile.ZipFile(target, "w") as archive:
        archive.writestr("mimetype", "application/epub+zip")
        archive.writestr("META-INF/container.xml", container)
        archive.writestr("OEBPS/content.opf", opf)
        archive.writestr("OEBPS/chapter1.xhtml", chapter)
    return target


def create_sample_pptx(target: Path) -> Path:
    presentation = Presentation()

    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Introduction"
    body = slide.placeholders[1].text_frame
    body.text = "Overview paragraph"
    bullet = body.add_paragraph()
    bullet.text = "First bullet"
    bullet.level = 1

    second_slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    textbox = second_slide.shapes.add_textbox(
        left=1000000,
        top=1200000,
        width=6000000,
        height=3000000,
    )
    frame = textbox.text_frame
    frame.text = "Body-only slide"
    frame.add_paragraph().text = "Follow-up point"

    presentation.save(target)
    return target


def assert_contains(text: str, expected: str):
    if expected not in text:
        raise AssertionError(f"Expected to find {expected!r}")


def test_single_file_conversion():
    with tempfile.TemporaryDirectory() as temp_dir:
        temp_path = Path(temp_dir)
        pdf_path = create_sample_pdf(temp_path / "book.pdf")
        output_dir = temp_path / "out"

        converter = DocumentToMarkdownConverter()
        results = converter.convert(pdf_path, output_dir)

        if len(results) != 1:
            raise AssertionError(f"Expected 1 Markdown file, got {len(results)}")

        markdown_path = results[0]
        if not markdown_path.exists():
            raise AssertionError("Markdown file was not created")

        content = markdown_path.read_text(encoding="utf-8")
        assert_contains(content, "# Sample PDF Book")
        assert_contains(content, "## Getting Started")
        assert_contains(
            content,
            "This is the first paragraph on page one. It should appear in the Markdown output.",
        )
        assert_contains(content, "1. First numbered item")


def test_mixed_directory_conversion():
    with tempfile.TemporaryDirectory() as temp_dir:
        temp_path = Path(temp_dir)
        input_dir = temp_path / "input"
        input_dir.mkdir()
        create_sample_pdf(input_dir / "book.pdf")
        create_sample_docx(input_dir / "book.docx")
        create_sample_epub(input_dir / "reader.epub")
        output_dir = temp_path / "out"

        converter = DocumentToMarkdownConverter()
        results = converter.convert(input_dir, output_dir)

        if len(results) != 3:
            raise AssertionError(f"Expected 3 Markdown files, got {len(results)}")

        output_names = {path.name for path in results}
        if "book.md" not in output_names or "reader.md" not in output_names:
            raise AssertionError(f"Unexpected output names: {sorted(output_names)}")
        collision_names = {
            name for name in output_names if name.startswith("book_") and name.endswith(".md")
        }
        if len(collision_names) != 1:
            raise AssertionError(f"Unexpected output names: {sorted(output_names)}")

        combined = "\n".join(path.read_text(encoding="utf-8") for path in results)
        assert_contains(combined, "# Sample PDF Book")
        assert_contains(combined, "# Sample DOCX Book")
        assert_contains(combined, "# Sample EPUB Book")
        assert_contains(combined, "DOCX list item")
        assert_contains(combined, "EPUB list item")


def test_powerpoint_conversion():
    with tempfile.TemporaryDirectory() as temp_dir:
        temp_path = Path(temp_dir)
        pptx_path = create_sample_pptx(temp_path / "deck.pptx")
        output_dir = temp_path / "out"

        converter = DocumentToMarkdownConverter()
        results = converter.convert(pptx_path, output_dir)

        if len(results) != 1:
            raise AssertionError(f"Expected 1 Markdown file, got {len(results)}")

        content = results[0].read_text(encoding="utf-8")
        assert_contains(content, "# Introduction")
        assert_contains(content, "## 01. Introduction")
        assert_contains(content, "Overview paragraph")
        assert_contains(content, "- First bullet")


def test_mixed_directory_includes_powerpoint():
    with tempfile.TemporaryDirectory() as temp_dir:
        temp_path = Path(temp_dir)
        input_dir = temp_path / "input"
        input_dir.mkdir()
        create_sample_docx(input_dir / "brief.docx")
        create_sample_pptx(input_dir / "brief.pptx")

        results = DocumentToMarkdownConverter().convert(input_dir, temp_path / "out")
        output_names = {path.name for path in results}
        if output_names != {"brief.md", "brief_pptx.md"}:
            raise AssertionError(f"Unexpected output names: {sorted(output_names)}")


def main():
    tests = [
        ("Single File Conversion", test_single_file_conversion),
        ("Mixed Directory Conversion", test_mixed_directory_conversion),
        ("PowerPoint Conversion", test_powerpoint_conversion),
        ("Mixed Directory Includes PowerPoint", test_mixed_directory_includes_powerpoint),
    ]
    failures = 0

    for name, func in tests:
        try:
            func()
            print(f"OK {name}")
        except Exception as exc:
            failures += 1
            print(f"FAIL {name}: {exc}")

    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
